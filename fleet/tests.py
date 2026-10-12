import io
import tempfile

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from PIL import Image
from django.urls import reverse

from .forms import VehicleForm
from .models import Vehicle, VehiclePhoto

MEDIA_ROOT = tempfile.mkdtemp()


def _png_upload(name="mobil.png"):
    buffer = io.BytesIO()
    Image.new("RGB", (4, 4), "red").save(buffer, format="PNG")
    return SimpleUploadedFile(name, buffer.getvalue(), content_type="image/png")


def _photos(count=4):
    return [_png_upload(f"mobil-{i}.png") for i in range(count)]


class VehicleModelTests(TestCase):
    def test_create_vehicle_and_str(self):
        vehicle = Vehicle.objects.create(
            nopol="B 9021 XY",
            jenis=Vehicle.Jenis.TRUK,
            merek="Hino Dutro",
            tahun=2019,
            status=Vehicle.Status.SIAP_PAKAI,
            penanggung_jawab="Budi Santoso",
        )
        self.assertEqual(str(vehicle), "B 9021 XY")
        self.assertEqual(vehicle.get_jenis_display(), "Truk")
        self.assertEqual(vehicle.get_status_display(), "Siap Pakai")

    def test_nopol_must_be_unique(self):
        Vehicle.objects.create(
            nopol="B 9021 XY", jenis=Vehicle.Jenis.TRUK, merek="Hino Dutro",
            tahun=2019, status=Vehicle.Status.SIAP_PAKAI,
            penanggung_jawab="Budi Santoso",
        )
        with self.assertRaises(Exception):
            Vehicle.objects.create(
                nopol="B 9021 XY", jenis=Vehicle.Jenis.JEEP, merek="Jeep Wrangler",
                tahun=2020, status=Vehicle.Status.SIAP_PAKAI,
                penanggung_jawab="Siti Amalia",
            )

    def test_default_ordering_is_by_nopol(self):
        Vehicle.objects.create(
            nopol="B 9021 XY", jenis=Vehicle.Jenis.TRUK, merek="Hino Dutro",
            tahun=2019, status=Vehicle.Status.SIAP_PAKAI,
            penanggung_jawab="Budi Santoso",
        )
        Vehicle.objects.create(
            nopol="B 1234 AB", jenis=Vehicle.Jenis.MOBIL, merek="Toyota Avanza",
            tahun=2021, status=Vehicle.Status.SIAP_PAKAI,
            penanggung_jawab="Siti Amalia",
        )
        nopols = list(Vehicle.objects.values_list("nopol", flat=True))
        self.assertEqual(nopols, ["B 1234 AB", "B 9021 XY"])


class VehicleFormTests(TestCase):
    def test_valid_data_is_valid(self):
        form = VehicleForm(data={
            "nopol": "B 9021 XY",
            "kepemilikan": "organik",
            "jenis": Vehicle.Jenis.TRUK,
            "merek": "Hino Dutro",
            "tahun": 2019,
            "status": Vehicle.Status.SIAP_PAKAI,
            "penanggung_jawab": "Budi Santoso",
        }, files={"foto": _photos()})
        self.assertTrue(form.is_valid(), form.errors)

    def test_missing_nopol_is_invalid(self):
        form = VehicleForm(data={
            "nopol": "",
            "kepemilikan": "organik",
            "jenis": Vehicle.Jenis.TRUK,
            "merek": "Hino Dutro",
            "tahun": 2019,
            "status": Vehicle.Status.SIAP_PAKAI,
            "penanggung_jawab": "Budi Santoso",
        })
        self.assertFalse(form.is_valid())
        self.assertIn("nopol", form.errors)

    def test_widgets_have_tailwind_classes(self):
        form = VehicleForm()
        self.assertIn("rounded-md", form.fields["nopol"].widget.attrs["class"])
        self.assertIn("rounded-md", form.fields["no_chasis"].widget.attrs["class"])

    def test_only_golongan_b_table_fields(self):
        self.assertEqual(
            list(VehicleForm().fields),
            ["nopol", "merek", "no_chasis", "no_mesin", "kepemilikan", "foto"],
        )


class VehicleListViewTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser(username="teknisi1", password="rahasia-123")

    def test_redirects_when_not_authenticated(self):
        response = self.client.get(reverse("fleet:vehicle_list"))
        self.assertRedirects(
            response,
            f"{reverse('accounts:login')}?next={reverse('fleet:vehicle_list')}",
        )

    def test_lists_vehicles(self):
        Vehicle.objects.create(
            nopol="B 9021 XY", jenis=Vehicle.Jenis.TRUK, merek="Hino Dutro",
            tahun=2019, status=Vehicle.Status.SIAP_PAKAI,
            penanggung_jawab="Budi Santoso", no_chasis="MHFXX123", no_mesin="W04DTR456",
        )
        self.client.login(username="teknisi1", password="rahasia-123")
        response = self.client.get(reverse("fleet:vehicle_list"))
        self.assertEqual(response.status_code, 200)
        for header in ("No. Registrasi", "Merk/ Type", "No. Chasis", "No. Mesin"):
            self.assertContains(response, header)
        self.assertContains(response, "B 9021 XY")
        self.assertContains(response, "Hino Dutro")
        self.assertNotContains(response, "Km / Mil")
        self.assertContains(response, "MHFXX123")
        self.assertContains(response, "W04DTR456")

    def test_empty_state(self):
        self.client.login(username="teknisi1", password="rahasia-123")
        response = self.client.get(reverse("fleet:vehicle_list"))
        self.assertContains(response, "Belum ada data kendaraan.")

    def test_tambah_button_opens_modal_via_htmx(self):
        self.client.login(username="teknisi1", password="rahasia-123")
        response = self.client.get(reverse("fleet:vehicle_list"))
        self.assertContains(response, f'hx-get="{reverse("fleet:vehicle_create")}?group=B"')
        self.assertContains(response, 'id="vehicle-modal-body"')

    def test_sidebar_armada_link_points_to_list(self):
        self.client.login(username="teknisi1", password="rahasia-123")
        response = self.client.get(reverse("fleet:vehicle_list"))
        self.assertContains(response, 'href="/armada/"')

    def test_delete_confirm_modal_escapes_nopol(self):
        vehicle = Vehicle.objects.create(nopol="B 1' <b>X</b>", merek="Test")
        self.client.login(username="teknisi1", password="rahasia-123")
        response = self.client.get(reverse("fleet:vehicle_delete_confirm", args=[vehicle.uuid]))
        self.assertContains(response, "B 1&#x27; &lt;b&gt;X&lt;/b&gt;")
        self.assertNotContains(response, "<b>X</b>")


@override_settings(MEDIA_ROOT=MEDIA_ROOT)
class VehicleCreateViewTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser(username="teknisi1", password="rahasia-123")
        self.client.login(username="teknisi1", password="rahasia-123")

    def test_get_renders_empty_form(self):
        response = self.client.get(reverse("fleet:vehicle_create"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Tambah Kendaraan")

    def test_post_valid_data_creates_vehicle_and_redirects(self):
        response = self.client.post(reverse("fleet:vehicle_create"), {
            "nopol": "B 1234 AB",
            "kepemilikan": "organik",
            "jenis": Vehicle.Jenis.MOBIL,
            "merek": "Toyota Avanza",
            "tahun": 2021,
            "status": Vehicle.Status.SIAP_PAKAI,
            "penanggung_jawab": "Siti Amalia",
            "foto": _photos(),
        })
        self.assertRedirects(response, reverse("fleet:vehicle_list"))
        self.assertEqual(Vehicle.objects.count(), 1)
        self.assertEqual(Vehicle.objects.first().nopol, "B 1234 AB")

    def test_post_missing_required_field_reshows_form_with_error(self):
        response = self.client.post(reverse("fleet:vehicle_create"), {
            "nopol": "",
            "kepemilikan": "organik",
            "jenis": Vehicle.Jenis.MOBIL,
            "merek": "Toyota Avanza",
            "tahun": 2021,
            "status": Vehicle.Status.SIAP_PAKAI,
            "penanggung_jawab": "Siti Amalia",
        })
        self.assertEqual(response.status_code, 200)
        self.assertEqual(Vehicle.objects.count(), 0)
        self.assertContains(response, "Bidang ini tidak boleh kosong.")

    def test_htmx_get_returns_form_partial_without_full_page(self):
        response = self.client.get(reverse("fleet:vehicle_create"), HTTP_HX_REQUEST="true")
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Tambah Kendaraan")
        self.assertNotContains(response, "<html")

    def test_htmx_post_valid_data_returns_hx_redirect(self):
        response = self.client.post(
            reverse("fleet:vehicle_create"),
            {
                "nopol": "B 1234 AB",
                "kepemilikan": "organik",
                "jenis": Vehicle.Jenis.MOBIL,
                "merek": "Toyota Avanza",
                "tahun": 2021,
                "status": Vehicle.Status.SIAP_PAKAI,
                "penanggung_jawab": "Siti Amalia",
                "foto": _photos(),
            },
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 204)
        self.assertEqual(response["HX-Redirect"], reverse("fleet:vehicle_list"))
        self.assertEqual(Vehicle.objects.count(), 1)

    def test_post_duplicate_nopol_reshows_form_with_error(self):
        Vehicle.objects.create(
            nopol="B 1234 AB", jenis=Vehicle.Jenis.MOBIL, merek="Toyota Avanza",
            tahun=2021, status=Vehicle.Status.SIAP_PAKAI,
            penanggung_jawab="Siti Amalia",
        )
        response = self.client.post(reverse("fleet:vehicle_create"), {
            "nopol": "B 1234 AB",
            "kepemilikan": "organik",
            "jenis": Vehicle.Jenis.MOBIL,
            "merek": "Toyota Avanza 2",
            "tahun": 2022,
            "status": Vehicle.Status.SIAP_PAKAI,
            "penanggung_jawab": "Budi Santoso",
        })
        self.assertEqual(response.status_code, 200)
        self.assertEqual(Vehicle.objects.count(), 1)
        self.assertContains(response, "telah ada")


@override_settings(MEDIA_ROOT=MEDIA_ROOT)
class VehicleUpdateViewTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser(username="teknisi1", password="rahasia-123")
        self.client.login(username="teknisi1", password="rahasia-123")
        self.vehicle = Vehicle.objects.create(
            nopol="B 9021 XY", jenis=Vehicle.Jenis.TRUK, merek="Hino Dutro",
            tahun=2019, status=Vehicle.Status.SIAP_PAKAI,
            penanggung_jawab="Budi Santoso",
        )
        for upload in _photos():
            VehiclePhoto.objects.create(vehicle=self.vehicle, image=upload)

    def test_get_renders_prefilled_form(self):
        response = self.client.get(reverse("fleet:vehicle_update", args=[self.vehicle.uuid]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Edit Kendaraan")
        self.assertContains(response, 'value="B 9021 XY"')

    def test_htmx_get_returns_form_partial_without_full_page(self):
        response = self.client.get(
            reverse("fleet:vehicle_update", args=[self.vehicle.uuid]), HTTP_HX_REQUEST="true"
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Edit Kendaraan")
        self.assertNotContains(response, "<html")

    def test_htmx_post_valid_data_returns_hx_redirect(self):
        response = self.client.post(
            reverse("fleet:vehicle_update", args=[self.vehicle.uuid]),
            {
                "nopol": "B 9021 XY",
                "kepemilikan": "organik",
                "jenis": Vehicle.Jenis.TRUK,
                "merek": "Hino Dutro",
                "tahun": 2019,
                "status": Vehicle.Status.DALAM_PERBAIKAN,
                "penanggung_jawab": "Budi Santoso",
            },
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 204)
        self.assertEqual(response["HX-Redirect"], reverse("fleet:vehicle_list"))

    def test_post_updates_vehicle_and_redirects(self):
        response = self.client.post(
            reverse("fleet:vehicle_update", args=[self.vehicle.uuid]),
            {
                "nopol": "B 9021 XY",
                "kepemilikan": "organik",
                "jenis": Vehicle.Jenis.TRUK,
                "merek": "Hino Dutro",
                "tahun": 2019,
                "status": Vehicle.Status.DALAM_PERBAIKAN,
                "penanggung_jawab": "Budi Santoso",
                "no_chasis": "MHFXX123",
            },
        )
        self.assertRedirects(response, reverse("fleet:vehicle_list"))
        self.vehicle.refresh_from_db()
        # Stored in capitals whatever was typed.
        self.assertEqual(self.vehicle.merek, "HINO DUTRO")
        self.assertEqual(self.vehicle.no_chasis, "MHFXX123")

    def test_get_nonexistent_vehicle_returns_404(self):
        response = self.client.get(reverse("fleet:vehicle_update", args=["00000000-0000-0000-0000-000000009999"]))
        self.assertEqual(response.status_code, 404)


@override_settings(MEDIA_ROOT=MEDIA_ROOT)
class VehicleDeleteViewTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser(username="teknisi1", password="rahasia-123")
        self.client.login(username="teknisi1", password="rahasia-123")
        self.vehicle = Vehicle.objects.create(
            nopol="B 9021 XY", jenis=Vehicle.Jenis.TRUK, merek="Hino Dutro",
            tahun=2019, status=Vehicle.Status.SIAP_PAKAI,
            penanggung_jawab="Budi Santoso",
        )

    def test_post_deletes_vehicle_and_redirects(self):
        response = self.client.post(reverse("fleet:vehicle_delete", args=[self.vehicle.uuid]))
        self.assertRedirects(response, reverse("fleet:vehicle_list"))
        self.assertEqual(Vehicle.objects.count(), 0)

    def test_get_is_not_allowed(self):
        response = self.client.get(reverse("fleet:vehicle_delete", args=[self.vehicle.uuid]))
        self.assertEqual(response.status_code, 405)
        self.assertEqual(Vehicle.objects.count(), 1)

    def test_list_hapus_opens_confirm_modal_instead_of_js_confirm(self):
        response = self.client.get(reverse("fleet:vehicle_list"))
        url = reverse("fleet:vehicle_delete_confirm", args=[self.vehicle.uuid])
        self.assertContains(response, f'hx-get="{url}"', count=2)
        self.assertNotContains(response, "confirm(")

    def test_confirm_modal_posts_to_delete(self):
        response = self.client.get(reverse("fleet:vehicle_delete_confirm", args=[self.vehicle.uuid]))
        self.assertContains(response, "Hapus Kendaraan")
        self.assertContains(response, "B 9021 XY")
        self.assertContains(response, f'action="{reverse("fleet:vehicle_delete", args=[self.vehicle.uuid])}"')
        self.assertEqual(Vehicle.objects.count(), 1)

    def test_delete_removes_photo_files(self):
        foto = VehiclePhoto.objects.create(vehicle=self.vehicle, image=_png_upload())
        storage, name = foto.image.storage, foto.image.name
        self.assertTrue(storage.exists(name))
        self.client.post(reverse("fleet:vehicle_delete", args=[self.vehicle.uuid]))
        self.assertFalse(storage.exists(name))


class VehicleQrCodeTests(TestCase):
    def test_qr_text_lists_vehicle_identity(self):
        vehicle = Vehicle(nopol="B 9021 XY", merek="Hino Dutro", no_chasis="MHFXX123", no_mesin="W04DTR456")
        self.assertEqual(
            vehicle.qr_text(),
            "No. Registrasi: B 9021 XY\nMerk/Type: Hino Dutro\nNo. Chasis: MHFXX123\nNo. Mesin: W04DTR456",
        )

    def test_qr_png_has_logo_in_centre(self):
        from .qr import qr_png_with_logo
        image = Image.open(io.BytesIO(qr_png_with_logo("No. Registrasi: B 9021 XY")))
        self.assertEqual(image.format, "PNG")
        centre = image.getpixel((image.width // 2, image.height // 2))
        corner = image.getpixel((5, 5))
        self.assertEqual(corner, (255, 255, 255))
        self.assertNotIn(centre, [(0, 0, 0), (255, 255, 255)])

    def test_qr_png_label_adds_strip_below(self):
        from .qr import qr_png_with_logo
        plain = Image.open(io.BytesIO(qr_png_with_logo("No. Registrasi: B 9021 XY")))
        labelled = Image.open(io.BytesIO(qr_png_with_logo("No. Registrasi: B 9021 XY", label="B 9021 XY")))
        self.assertEqual(labelled.width, plain.width)
        self.assertGreater(labelled.height, plain.height)

    def test_qr_view_requires_login_and_serves_png(self):
        vehicle = Vehicle.objects.create(nopol="B 9021 XY", merek="Hino Dutro")
        url = reverse("fleet:vehicle_qr", args=[vehicle.uuid])
        self.assertEqual(self.client.get(url).status_code, 302)
        User.objects.create_superuser(username="teknisi1", password="rahasia-123")
        self.client.login(username="teknisi1", password="rahasia-123")
        response = self.client.get(url)
        self.assertEqual(response["Content-Type"], "image/png")
        self.assertEqual(response["Cache-Control"], "private, no-cache")
        self.assertNotIn("Content-Disposition", response)
        response = self.client.get(url + "?unduh=1")
        self.assertEqual(response["Content-Disposition"], 'attachment; filename="qr-b-9021-xy.png"')

    def test_list_shows_qr_code_and_foto_columns(self):
        User.objects.create_superuser(username="teknisi1", password="rahasia-123")
        Vehicle.objects.create(nopol="B 9021 XY", merek="Hino Dutro")
        self.client.login(username="teknisi1", password="rahasia-123")
        response = self.client.get(reverse("fleet:vehicle_list"))
        self.assertContains(response, "QR Code</th>")
        self.assertContains(response, "Foto Kendaraan</th>")
        vehicle = Vehicle.objects.get()
        # The list only shows an icon button; the QR image lives in the modal.
        self.assertContains(response, f'hx-get="{reverse("fleet:vehicle_qr_modal", args=[vehicle.uuid])}"', count=2)
        self.assertNotContains(response, 'alt="QR Code B 9021 XY"')

        response = self.client.get(reverse("fleet:vehicle_qr_modal", args=[vehicle.uuid]))
        self.assertContains(
            response, f'src="{reverse("fleet:vehicle_qr", args=[vehicle.uuid])}?v={vehicle.qr_version()}"'
        )
        # Without download, the page loading bar treats the link as a
        # navigation and stays stuck because the page never unloads.
        self.assertContains(response, 'download="qr-b-9021-xy.png"')

    def test_qr_version_changes_when_no_mesin_changes(self):
        vehicle = Vehicle(nopol="B 9021 XY", merek="Hino Dutro")
        before = vehicle.qr_version()
        vehicle.no_mesin = "W04DTR456"
        self.assertNotEqual(before, vehicle.qr_version())
        self.assertIn("No. Mesin: W04DTR456", vehicle.qr_text())


@override_settings(MEDIA_ROOT=MEDIA_ROOT)
class VehicleFotoUploadTests(TestCase):
    def setUp(self):
        User.objects.create_superuser(username="teknisi1", password="rahasia-123")
        self.client.login(username="teknisi1", password="rahasia-123")

    def _create(self, photos):
        return self.client.post(reverse("fleet:vehicle_create"), {
            "nopol": "B 1234 AB",
            "kepemilikan": "organik",
            "merek": "Toyota Avanza",
            "foto": photos,
        })

    def test_create_with_four_fotos_saves_them_and_lists_them(self):
        response = self._create(_photos(4))
        self.assertRedirects(response, reverse("fleet:vehicle_list"))
        vehicle = Vehicle.objects.get()
        self.assertEqual(vehicle.fotos.count(), 4)
        self.assertTrue(all(f.image.name.startswith("kendaraan/") for f in vehicle.fotos.all()))
        response = self.client.get(reverse("fleet:vehicle_list"))
        self.assertContains(response, f'hx-get="{reverse("fleet:vehicle_photos", args=[vehicle.uuid])}"', count=2)
        self.assertContains(response, "Lihat (4)")
        response = self.client.get(reverse("fleet:vehicle_photos", args=[vehicle.uuid]))
        self.assertContains(response, "Foto Kendaraan B 1234 AB")
        for foto in vehicle.fotos.all():
            self.assertContains(response, f'src="{foto.image.url}"')

    def test_uploaded_fotos_are_compressed_to_jpeg(self):
        buffer = io.BytesIO()
        Image.effect_noise((4000, 3000), 80).convert("RGB").save(buffer, format="PNG")
        big = buffer.getvalue()
        uploads = [SimpleUploadedFile(f"besar-{i}.png", big, content_type="image/png") for i in range(4)]
        self.assertRedirects(self._create(uploads), reverse("fleet:vehicle_list"))
        foto = Vehicle.objects.get().fotos.first()
        self.assertTrue(foto.image.name.endswith(".jpg"))
        self.assertLess(foto.image.size, len(big))
        with Image.open(foto.image) as img:
            self.assertEqual(img.format, "JPEG")
            self.assertEqual(max(img.size), 1600)

    def test_photos_modal_requires_login(self):
        self._create(_photos(4))
        vehicle = Vehicle.objects.get()
        self.client.logout()
        response = self.client.get(reverse("fleet:vehicle_photos", args=[vehicle.uuid]))
        self.assertEqual(response.status_code, 302)

    def test_create_with_fewer_than_four_fotos_is_rejected(self):
        response = self._create(_photos(3))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(Vehicle.objects.count(), 0)
        self.assertContains(response, "Unggah minimal 4 foto kendaraan (baru 3).")

    def test_create_without_fotos_is_rejected(self):
        response = self._create([])
        self.assertEqual(Vehicle.objects.count(), 0)
        self.assertContains(response, "Unggah minimal 4 foto kendaraan (baru 0).")

    def test_edit_with_new_fotos_replaces_all_old_ones(self):
        self._create(_photos(4))
        vehicle = Vehicle.objects.get()
        old_ids = set(vehicle.fotos.values_list("id", flat=True))
        response = self.client.post(reverse("fleet:vehicle_update", args=[vehicle.uuid]), {
            "kepemilikan": "organik", "nopol": "B 1234 AB", "merek": "Toyota Avanza", "foto": _photos(4),
        })
        self.assertRedirects(response, reverse("fleet:vehicle_list"))
        self.assertEqual(vehicle.fotos.count(), 4)
        self.assertFalse(old_ids & set(vehicle.fotos.values_list("id", flat=True)))

    def test_edit_with_too_few_new_fotos_keeps_old_ones(self):
        self._create(_photos(4))
        vehicle = Vehicle.objects.get()
        response = self.client.post(reverse("fleet:vehicle_update", args=[vehicle.uuid]), {
            "kepemilikan": "organik", "nopol": "B 1234 AB", "merek": "Toyota Avanza", "foto": _photos(2),
        })
        self.assertEqual(response.status_code, 200)
        self.assertEqual(vehicle.fotos.count(), 4)

    def test_form_posts_as_multipart_for_htmx(self):
        response = self.client.get(reverse("fleet:vehicle_create"), HTTP_HX_REQUEST="true")
        self.assertContains(response, 'hx-encoding="multipart/form-data"')


class AdminRegistrationTests(TestCase):
    def test_every_project_model_is_registered_in_admin(self):
        from django.apps import apps
        from django.contrib import admin

        for app_label in ("accounts", "fleet"):
            for model in apps.get_app_config(app_label).get_models():
                self.assertTrue(admin.site.is_registered(model), f"{model.__name__} not in admin")

    def test_admin_changelists_render(self):
        User.objects.create_superuser(username="admin", password="rahasia-123")
        self.client.login(username="admin", password="rahasia-123")
        Vehicle.objects.create(nopol="B 9021 XY", merek="Hino Dutro")
        for name in ("fleet_vehicle", "fleet_vehiclephoto", "accounts_laporan", "accounts_revokedtoken"):
            response = self.client.get(reverse(f"admin:{name}_changelist"))
            self.assertEqual(response.status_code, 200, name)


@override_settings(MEDIA_ROOT=MEDIA_ROOT)
class VehicleGroupKendaraanTests(TestCase):
    def setUp(self):
        User.objects.create_superuser(username="teknisi1", password="rahasia-123")
        self.client.login(username="teknisi1", password="rahasia-123")

    def test_defaults_to_golongan_b(self):
        vehicle = Vehicle.objects.create(nopol="B 9021 XY", merek="Hino Dutro")
        self.assertEqual(vehicle.group_kendaraan, "B")
        self.assertEqual(vehicle.get_group_kendaraan_display(), "Golongan B")

    def test_create_from_golongan_b_tab_stores_group(self):
        response = self.client.get(reverse("fleet:vehicle_list"))
        self.assertContains(response, f'hx-get="{reverse("fleet:vehicle_create")}?group=B"')
        self.assertContains(response, "GOLONGAN B")
        form = self.client.get(reverse("fleet:vehicle_create") + "?group=B", HTTP_HX_REQUEST="true")
        self.assertContains(form, f'hx-post="{reverse("fleet:vehicle_create")}?group=B"')
        self.client.post(reverse("fleet:vehicle_create") + "?group=B", {
            "kepemilikan": "organik", "nopol": "B 1234 AB", "merek": "Toyota Avanza", "foto": _photos(),
        })
        self.assertEqual(Vehicle.objects.get().group_kendaraan, "B")

    def test_unknown_group_falls_back_to_golongan_b(self):
        self.client.post(reverse("fleet:vehicle_create") + "?group=ZZZ", {
            "kepemilikan": "organik", "nopol": "B 1234 AB", "merek": "Toyota Avanza", "foto": _photos(),
        })
        self.assertEqual(Vehicle.objects.get().group_kendaraan, "B")

    def test_list_only_shows_vehicles_of_the_tab_group(self):
        Vehicle.objects.create(nopol="B 1111 AA", merek="A")
        Vehicle.objects.filter(nopol="B 1111 AA").update(group_kendaraan="X")
        Vehicle.objects.create(nopol="B 2222 BB", merek="B")
        response = self.client.get(reverse("fleet:vehicle_list"))
        self.assertContains(response, "B 2222 BB")
        self.assertNotContains(response, "B 1111 AA")

    def test_admin_shows_and_filters_group(self):
        User.objects.create_superuser(username="admin", password="rahasia-123")
        self.client.login(username="admin", password="rahasia-123")
        Vehicle.objects.create(nopol="B 9021 XY", merek="Hino Dutro")
        response = self.client.get(reverse("admin:fleet_vehicle_changelist"))
        self.assertContains(response, "Group Kendaraan")
        self.assertContains(response, "Golongan B")


class VehicleLookupTests(TestCase):
    def setUp(self):
        User.objects.create_superuser(username="teknisi1", password="rahasia-123")
        Vehicle.objects.create(nopol="B 9021 XY", merek="Hino Dutro", no_chasis="MHFXX123", no_mesin="W04DTR456")

    def test_requires_login(self):
        response = self.client.get(reverse("fleet:vehicle_lookup"), {"nopol": "B 9021 XY"})
        self.assertEqual(response.status_code, 302)

    def test_returns_fields_shown_after_scan(self):
        self.client.login(username="teknisi1", password="rahasia-123")
        response = self.client.get(reverse("fleet:vehicle_lookup"), {"nopol": " B 9021 XY "})
        self.assertEqual(response.json(), {
            "uuid": str(Vehicle.objects.get().uuid),
            "nopol": "B 9021 XY",
            "merek": "Hino Dutro",
            "no_chasis": "MHFXX123",
            "no_mesin": "W04DTR456",
            "group_kendaraan": "Golongan B",
            "group_kode": "B",
        })

    def test_unknown_nopol_returns_404(self):
        self.client.login(username="teknisi1", password="rahasia-123")
        response = self.client.get(reverse("fleet:vehicle_lookup"), {"nopol": "Z 0000 ZZ"})
        self.assertEqual(response.status_code, 404)
        self.assertEqual(response.json(), {"error": "Kendaraan tidak ditemukan."})



@override_settings(MEDIA_ROOT=MEDIA_ROOT)
class KendaraanBermotorTabTests(TestCase):
    def setUp(self):
        User.objects.create_superuser(username="teknisi1", password="rahasia-123")
        self.client.login(username="teknisi1", password="rahasia-123")
        Vehicle.objects.create(nopol="B 1111 AA", merek="Truk")
        Vehicle.objects.create(nopol="B 2222 BB", merek="Motor", group_kendaraan="bermotor")

    def test_armada_shows_both_tabs(self):
        response = self.client.get(reverse("fleet:vehicle_list"))
        self.assertContains(response, "GOLONGAN B")
        self.assertContains(response, "MOTOR\n")
        self.assertContains(response, f'href="{reverse("fleet:vehicle_list")}?group=bermotor"')

    def test_bermotor_tab_lists_only_its_vehicles_and_adds_to_it(self):
        response = self.client.get(reverse("fleet:vehicle_list"), {"group": "bermotor"})
        self.assertContains(response, "B 2222 BB")
        self.assertNotContains(response, "B 1111 AA")
        self.assertContains(response, f'hx-get="{reverse("fleet:vehicle_create")}?group=bermotor"')
        response = self.client.post(reverse("fleet:vehicle_create") + "?group=bermotor", {
            "kepemilikan": "organik", "nopol": "B 3333 CC", "merek": "Honda", "foto": _photos(),
        })
        self.assertRedirects(response, reverse("fleet:vehicle_list") + "?group=bermotor")
        self.assertEqual(Vehicle.objects.get(nopol="B 3333 CC").group_kendaraan, "bermotor")

    def test_bermotor_vehicle_is_not_accepted_by_golongan_b_inspection(self):
        from accounts.forms import PemeriksaanGolonganBForm
        motor = Vehicle.objects.get(nopol="B 2222 BB")
        form = PemeriksaanGolonganBForm(data={"vehicle": motor.uuid})
        form.is_valid()
        self.assertIn("vehicle", form.errors)



class VehicleFotoSlotsTests(TestCase):
    def setUp(self):
        User.objects.create_superuser(username="teknisi1", password="rahasia-123")
        self.client.login(username="teknisi1", password="rahasia-123")

    def test_create_form_has_four_required_photo_slots(self):
        response = self.client.get(reverse("fleet:vehicle_create"), HTTP_HX_REQUEST="true")
        content = response.content.decode()
        for nomor in range(1, 5):
            self.assertRegex(content, rf'id="foto_{nomor}" name="foto"[^>]*required')
        self.assertNotContains(response, 'id="foto_tambahan"')
        self.assertContains(response, "(4 foto)")

    def test_edit_form_slots_are_optional(self):
        vehicle = Vehicle.objects.create(nopol="B 1 A", merek="X")
        response = self.client.get(reverse("fleet:vehicle_update", args=[vehicle.uuid]), HTTP_HX_REQUEST="true")
        self.assertNotRegex(response.content.decode(), r'id="foto_1" name="foto"[^>]*required')



class VehicleSearchTests(TestCase):
    def setUp(self):
        User.objects.create_superuser(username="teknisi1", password="rahasia-123")
        self.client.login(username="teknisi1", password="rahasia-123")
        Vehicle.objects.create(nopol="B 1111 AA", merek="Hino Dutro", no_chasis="CH-AAA", no_mesin="MS-111")
        Vehicle.objects.create(nopol="B 2222 BB", merek="Honda Beat", group_kendaraan="bermotor", no_mesin="MS-222")

    def test_list_has_cari_button(self):
        response = self.client.get(reverse("fleet:vehicle_list"))
        self.assertContains(response, f'hx-get="{reverse("fleet:vehicle_search")}"')
        self.assertContains(response, "Cari")

    def test_modal_shows_search_field_and_all_vehicles(self):
        response = self.client.get(reverse("fleet:vehicle_search"))
        self.assertContains(response, 'type="search" name="q"')
        self.assertContains(response, "B 1111 AA")
        self.assertContains(response, "B 2222 BB")
        self.assertContains(response, "Motor")

    def test_live_search_filters_rows(self):
        for q, found, missing in (("hino", "B 1111 AA", "B 2222 BB"), ("MS-222", "B 2222 BB", "B 1111 AA"),
                                  ("ch-aaa", "B 1111 AA", "B 2222 BB"), ("2222", "B 2222 BB", "B 1111 AA")):
            response = self.client.get(reverse("fleet:vehicle_search"), {"q": q, "hasil": "1"})
            self.assertContains(response, found)
            self.assertNotContains(response, missing)
            self.assertNotContains(response, 'type="search"')

    def test_no_match_message(self):
        response = self.client.get(reverse("fleet:vehicle_search"), {"q": "zzz", "hasil": "1"})
        self.assertContains(response, 'Tidak ada kendaraan yang cocok dengan')
        self.assertContains(response, 'zzz')



@override_settings(MEDIA_ROOT=MEDIA_ROOT)
class KendaraanDieselTabTests(TestCase):
    def setUp(self):
        User.objects.create_superuser(username="teknisi1", password="rahasia-123")
        self.client.login(username="teknisi1", password="rahasia-123")
        Vehicle.objects.create(nopol="B 1111 AA", merek="Hino")
        Vehicle.objects.create(nopol="B 3333 DD", merek="Isuzu Elf", group_kendaraan="diesel", no_mesin="4JB1")

    def test_diesel_tab_with_same_table(self):
        response = self.client.get(reverse("fleet:vehicle_list"))
        self.assertContains(response, "KENDARAAN DIESEL")
        self.assertContains(response, f'href="{reverse("fleet:vehicle_list")}?group=diesel"')
        response = self.client.get(reverse("fleet:vehicle_list"), {"group": "diesel"})
        for header in ("No. Registrasi", "Merk/ Type", "No. Chasis", "No. Mesin", "QR Code", "Foto Kendaraan", "Aksi"):
            self.assertContains(response, header)
        self.assertContains(response, "B 3333 DD")
        self.assertNotContains(response, "B 1111 AA")
        self.assertContains(response, f'hx-get="{reverse("fleet:vehicle_create")}?group=diesel"')

    def test_add_diesel_vehicle(self):
        response = self.client.post(reverse("fleet:vehicle_create") + "?group=diesel", {
            "kepemilikan": "organik", "nopol": "B 4444 DD", "merek": "Mitsubishi Canter", "foto": _photos(),
        })
        self.assertRedirects(response, reverse("fleet:vehicle_list") + "?group=diesel")
        self.assertEqual(Vehicle.objects.get(nopol="B 4444 DD").group_kendaraan, "diesel")


class VehicleSearchHistoryModalTests(TestCase):
    def setUp(self):
        User.objects.create_superuser(username="teknisi1", password="rahasia-123")
        self.client.login(username="teknisi1", password="rahasia-123")
        self.vehicle = Vehicle.objects.create(nopol="B 1111 AA", merek="Hino Dutro", no_chasis="CH-1")

    def test_rows_open_the_history_modal(self):
        response = self.client.get(reverse("fleet:vehicle_search"))
        self.assertContains(response, f'hx-get="{reverse("accounts:history_kendaraan_modal")}?nopol=B%201111%20AA"')
        self.assertContains(response, "Klik baris kendaraan")

    def test_history_modal_has_the_history_tabs(self):
        response = self.client.get(reverse("accounts:history_kendaraan_modal"), {"nopol": "B 1111 AA"})
        for text in ("History B 1111 AA", "PROFILE", "STATUS KENDARAAN", "RIWAYAT PEMERIKSAAN", "RIWAYAT PERBAIKAN",
                     "x-data=\"{ tab: 'profile' }\"", "CH-1", "Belum ada riwayat perbaikan",
                     f'hx-get="{reverse("fleet:vehicle_search")}"'):
            self.assertContains(response, text)


class JadwalServisTests(TestCase):
    def setUp(self):
        User.objects.create_superuser(username="teknisi1", password="rahasia-123")
        self.client.login(username="teknisi1", password="rahasia-123")
        self.vehicle = Vehicle.objects.create(nopol="B 9021 XY", merek="Hino Dutro")

    def test_sidebar_links_to_both_pages(self):
        response = self.client.get(reverse("operasional:jadwal_list"))
        self.assertContains(response, f'<a href="{reverse("operasional:jadwal_list")}" class=')
        self.assertContains(response, f'<a href="{reverse("operasional:servis_list")}" class=')

    def test_jadwal_table_crud_and_overdue_badge(self):
        from datetime import date, timedelta
        from fleet.models import Jadwal
        kemarin = (date.today() - timedelta(days=1)).isoformat()
        response = self.client.post(reverse("operasional:jadwal_create"), {
            "vehicle": self.vehicle.pk, "tanggal": kemarin, "kegiatan": "servis", "status": "terjadwal", "keterangan": "Ganti oli",
        }, HTTP_HX_REQUEST="true")
        self.assertEqual(response["HX-Redirect"], reverse("operasional:jadwal_list"))
        page = self.client.get(reverse("operasional:jadwal_list"))
        for text in ("Tanggal", "Kendaraan", "Kegiatan", "Status", "B 9021 XY", "Servis", "Ganti oli"):
            self.assertContains(page, text)
        self.assertContains(page, ">Terlambat</span>")
        jadwal = Jadwal.objects.get()
        self.client.post(reverse("operasional:jadwal_update", args=[jadwal.uuid]), {
            "vehicle": self.vehicle.pk, "tanggal": kemarin, "kegiatan": "servis", "status": "selesai",
        })
        page = self.client.get(reverse("operasional:jadwal_list"))
        self.assertNotContains(page, ">Terlambat</span>")
        self.assertContains(page, ">Selesai</span>")
        self.assertContains(self.client.get(reverse("operasional:jadwal_delete", args=[jadwal.uuid])), "Yakin ingin menghapus")
        self.client.post(reverse("operasional:jadwal_delete", args=[jadwal.uuid]))
        self.assertFalse(Jadwal.objects.exists())

    def test_servis_table_has_no_biaya(self):
        self.client.post(reverse("operasional:servis_create"), {
            "vehicle": self.vehicle.pk, "tanggal": "2026-09-01", "jenis_servis": "Ganti oli", "bengkel": "Bengkel 202",
        })
        page = self.client.get(reverse("operasional:servis_list"))
        for text in ("Jenis Servis", "Bengkel / Tempat", "Ganti oli", "Bengkel 202"):
            self.assertContains(page, text)
        for text in ("Biaya", "Rp "):
            self.assertNotContains(page, text)
        self.assertNotContains(self.client.get(reverse("operasional:servis_create")), 'name="biaya"')

    def test_servis_exports_to_excel_instead_of_tambah(self):
        from io import BytesIO

        from openpyxl import load_workbook
        self.client.post(reverse("operasional:servis_create"), {
            "vehicle": self.vehicle.pk, "tanggal": "2026-09-01", "jenis_servis": "Ganti oli", "bengkel": "Bengkel 202",
        })
        page = self.client.get(reverse("operasional:servis_list"))
        self.assertContains(page, "Export ke Excel")
        self.assertNotContains(page, "Tambah Servis")
        response = self.client.get(reverse("operasional:servis_export"))
        self.assertIn(".xlsx", response["Content-Disposition"])
        rows = list(load_workbook(BytesIO(response.content)).active.values)
        self.assertEqual(rows[3][:7], ("No", "Tanggal", "No. Registrasi", "Merk/Type", "Satuan", "Form No", "Jenis Servis"))
        self.assertEqual(rows[4][:8], (1, "01-09-2026", self.vehicle.nopol, self.vehicle.merek, None, None, "Ganti oli", "Bengkel 202"))
        # A typed "=..." stays text, not a formula Excel would run.
        self.client.post(reverse("operasional:servis_create"), {
            "vehicle": self.vehicle.pk, "tanggal": "2026-09-02", "jenis_servis": "=1+1", "bengkel": "B",
        })
        sel = load_workbook(BytesIO(self.client.get(reverse("operasional:servis_export")).content)).active["G5"]
        self.assertEqual((sel.value, sel.data_type), ("=1+1", "s"))
        from accounts import roles
        self.assertEqual(roles.menus_of_view("operasional:servis_export"), {"servis"})
        self.assertContains(page, ">Satuan</th>")

    def test_repairs_from_form_perbaikan_are_listed_too(self):
        from accounts.models import (
            Laporan, PemeriksaanGolonganB, PemeriksaanGolonganBItem, PerbaikanGolonganB, Satuan,
        )
        teknisi = Laporan.objects.create(
            nip="1", nama="Budi Teknisi", pangkat="P", jabatan="J", satuan=Satuan.objects.create(nama="S"), foto="x.jpg",
        )
        periksa = PemeriksaanGolonganB.objects.create(
            pemeriksa=teknisi, vehicle=self.vehicle, pembagian_kelas="A", eselon_reparasi="1",
            macam_pemeriksaan="tehnis", buku_jalan="ada",
        )
        item = PemeriksaanGolonganBItem.objects.create(pemeriksaan=periksa, nomor=7, nama_alat="Kopling", kondisi="RB")
        PerbaikanGolonganB.objects.create(teknisi=teknisi, item=item, kondisi="B", keterangan="ganti plat")
        # A repair that left the part Rusak is not a service.
        rusak = PemeriksaanGolonganBItem.objects.create(pemeriksaan=periksa, nomor=8, nama_alat="Gardan", kondisi="RB")
        PerbaikanGolonganB.objects.create(teknisi=teknisi, item=rusak, kondisi="RB", keterangan="tidak bisa")
        self.client.post(reverse("operasional:servis_create"), {
            "vehicle": self.vehicle.pk, "tanggal": "2020-01-01", "jenis_servis": "Ganti oli", "bengkel": "Bengkel 202",
        })
        page = self.client.get(reverse("operasional:servis_list"))
        for text in ("Perbaikan Kopling", '<td class="py-2 px-4 tabular-nums">7</td>', "Budi Teknisi", "Kondisi: Baik · ganti plat", "Form Perbaikan",
                     "Ganti oli", "Menampilkan 1–2 dari 2 data"):
            self.assertContains(page, text)
        # Newest first, and only the service recorded here can be edited.
        content = page.content.decode()
        self.assertLess(content.index("Perbaikan Kopling"), content.index("Ganti oli"))
        self.assertNotIn("Gardan", content)
        # The filter: dates, part of a No. Registrasi, and what the repair left the part as.
        lain = PemeriksaanGolonganBItem.objects.create(pemeriksaan=periksa, nomor=9, nama_alat="Aki", kondisi="RR")
        PerbaikanGolonganB.objects.create(teknisi=teknisi, item=lain, kondisi="RR", keterangan="tunggu")
        url = reverse("operasional:servis_list")

        def jenis(**query):
            return [b["jenis_servis"] for b in self.client.get(url, query).context["servis_list"]]

        self.assertEqual(jenis(kondisi="B"), ["Perbaikan Kopling"])
        self.assertEqual(jenis(kondisi="RR"), ["Perbaikan Aki"])
        self.assertEqual(jenis(sampai="2020-12-31"), ["Ganti oli"])
        self.assertEqual(len(jenis(dari="2021-01-01")), 2)
        self.assertEqual(len(jenis(nopol=self.vehicle.nopol[:3].lower())), 3)
        self.assertEqual(jenis(nopol="tidak-ada"), [])
        self.assertEqual(len(jenis(kondisi="zz", dari="bukan-tanggal")), 3)
        from io import BytesIO

        from openpyxl import load_workbook
        saring = self.client.get(url, {"kondisi": "B"})
        self.assertContains(saring, "servis/export/?kondisi=B")
        rows = list(load_workbook(BytesIO(self.client.get(reverse("operasional:servis_export"), {"kondisi": "B"}).content)).active.values)
        self.assertEqual([r[6] for r in rows[4:]], ["Perbaikan Kopling"])
        self.assertEqual(content.count(">Edit</button>"), 1)

    def test_required_fields(self):
        response = self.client.post(reverse("operasional:servis_create"), {"vehicle": self.vehicle.pk})
        self.assertEqual(response.status_code, 200)
        from fleet.models import Servis
        self.assertFalse(Servis.objects.exists())

    def test_vehicle_with_servis_history_cannot_be_deleted(self):
        from fleet.models import Servis
        Servis.objects.create(vehicle=self.vehicle, tanggal="2026-09-01", jenis_servis="Tune up")
        self.client.post(reverse("fleet:vehicle_delete", args=[self.vehicle.uuid]))
        self.assertTrue(Vehicle.objects.filter(pk=self.vehicle.pk).exists())

    def test_tim_users_have_no_access(self):
        from django.contrib.auth.models import Group
        user = User.objects.create_user(username="t", password="rahasia-123")
        user.groups.add(Group.objects.get(name="Tim Perbaikan"))
        self.client.login(username="t", password="rahasia-123")
        self.assertRedirects(self.client.get(reverse("operasional:servis_list")), reverse("accounts:perbaikan"))


class VehicleFotoMaksimalTests(TestCase):
    def test_more_than_four_photos_is_refused(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        from fleet.forms import VehicleForm

        gif = (b"GIF89a\x01\x00\x01\x00\x80\x00\x00\x00\x00\x00\xff\xff\xff!\xf9\x04\x01\x00\x00\x00\x00"
               b",\x00\x00\x00\x00\x01\x00\x01\x00\x00\x02\x02D\x01\x00;")
        from django.utils.datastructures import MultiValueDict

        files = MultiValueDict({"foto": [SimpleUploadedFile(f"{i}.gif", gif, content_type="image/gif") for i in range(5)]})
        form = VehicleForm({"kepemilikan": "organik", "nopol": "B 1 X", "merek": "Hino"}, files)
        self.assertFalse(form.is_valid())
        self.assertIn("maksimal 4", form.errors["foto"][0])


class RanpurRantisTests(TestCase):
    """RANPUR and RANTIS: vehicle groups with their own tab and table, no inspection form yet."""

    def setUp(self):
        User.objects.create_superuser(username="root", password="rahasia-123")
        self.client.login(username="root", password="rahasia-123")

    def test_tabs_and_tables(self):
        Vehicle.objects.create(nopol="RP 01", merek="Anoa", group_kendaraan="ranpur")
        Vehicle.objects.create(nopol="RT 01", merek="Komodo", group_kendaraan="rantis")
        daftar = reverse("fleet:vehicle_list")
        halaman = self.client.get(daftar)
        for kode, label in (("ranpur", "RANPUR"), ("rantis", "RANTIS")):
            self.assertContains(halaman, f'<a href="{daftar}?group={kode}"')
            self.assertContains(halaman, label)
        ranpur = self.client.get(daftar, {"group": "ranpur"})
        self.assertContains(ranpur, "RP 01")
        self.assertNotContains(ranpur, "RT 01")
        self.assertContains(self.client.get(daftar, {"group": "rantis"}), "RT 01")

    def test_can_be_added_from_their_tab_and_form_tambah(self):
        self.assertEqual(self.client.get(reverse("fleet:vehicle_create") + "?group=rantis").status_code, 200)
        from accounts import golongan, roles
        self.assertIn(("ranpur", "RANPUR"), roles.FORM_TAMBAH)
        self.assertIn(("rantis", "RANTIS"), roles.FORM_TAMBAH)
        # No inspection form yet: scanning one says so instead of failing.
        self.assertIsNone(golongan.for_vehicle(Vehicle(group_kendaraan="ranpur")))
        self.assertEqual(self.client.get(reverse("accounts:pemeriksaan_checklist", args=["ranpur"])).status_code, 404)
