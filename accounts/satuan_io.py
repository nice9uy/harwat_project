"""Export and Import of Pengaturan > Satuan: the list as an Excel sheet, and
an Excel/CSV sheet read back in."""
import csv
import io

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import render
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from . import roles
from .models import Satuan

KEPALA = ["Nama Satuan", "Matra", "Keterangan"]
MATRA = dict(roles.MATRA)  # kode -> label
PANJANG_NAMA = Satuan._meta.get_field("nama").max_length


@login_required
def satuan_ekspor(request):
    wb = Workbook()
    ws = wb.active
    ws.title = "Satuan"
    ws.append(KEPALA)
    for sel in ws[1]:
        sel.font = Font(bold=True, color="FFFFFF")
        sel.fill = PatternFill("solid", fgColor="0B1F3A")
    for nama, matra, keterangan in Satuan.objects.order_by("matra", "nama").values_list("nama", "matra", "keterangan"):
        ws.append([nama, MATRA.get(matra, matra), keterangan])
    from .laporan import teks_saja
    teks_saja(ws)
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 60
    ws.freeze_panes = "A2"
    isi = io.BytesIO()
    wb.save(isi)
    response = HttpResponse(
        isi.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="satuan-{timezone.localtime():%Y%m%d}.xlsx"'
    return response


def _baris_berkas(berkas):
    """Rows of the uploaded sheet as lists of strings (.xlsx or .csv)."""
    if berkas.name.lower().endswith(".csv"):
        teks = berkas.read().decode("utf-8-sig")
        pemisah = ";" if teks.split("\n", 1)[0].count(";") > teks.split("\n", 1)[0].count(",") else ","
        return [[sel.strip() for sel in baris] for baris in csv.reader(io.StringIO(teks), delimiter=pemisah)]
    ws = load_workbook(berkas, read_only=True, data_only=True).active
    return [["" if sel is None else str(sel).strip() for sel in baris] for baris in ws.iter_rows(values_only=True)]


def _matra(teks):
    """A sheet's Matra cell ("Matra Laut", "laut", "Laut") as its code, or ''."""
    teks = teks.strip().lower().removeprefix("matra").strip()
    return teks if teks in MATRA else ""


def impor_baris(baris, matra_bawaan="darat"):
    """Add the units of ``baris`` (a header row, then name, matra and note).
    A name already there is kept, with its note and matra replaced when the
    sheet gives others. A new unit without a Matra cell gets ``matra_bawaan``.
    Returns the counts."""
    baris = [b for b in baris if any(b)]
    hasil = {"baru": 0, "diperbarui": 0, "sama": 0, "dilewati": 0}
    if not baris:
        return hasil
    kepala = [sel.lower() for sel in baris[0]]
    i_nama = next((i for i, k in enumerate(kepala) if k in ("nama satuan", "nama", "satuan")), None)
    if i_nama is None:
        # No header: the first column is the name, the second the note.
        i_nama, i_ket, i_matra = 0, 1, None
    else:
        i_ket = next((i for i, k in enumerate(kepala) if k == "keterangan"), None)
        i_matra = next((i for i, k in enumerate(kepala) if k == "matra"), None)
        baris = baris[1:]
    ada = {s.nama: s for s in Satuan.objects.all()}
    baru, ubah = {}, []
    for b in baris:
        nama = b[i_nama] if i_nama < len(b) else ""
        keterangan = b[i_ket] if i_ket is not None and i_ket < len(b) else ""
        matra = _matra(b[i_matra]) if i_matra is not None and i_matra < len(b) else ""
        if not nama or len(nama) > PANJANG_NAMA or nama in baru:
            hasil["dilewati"] += 1
        elif nama in ada:
            berubah = False
            if keterangan and keterangan != ada[nama].keterangan:
                ada[nama].keterangan, berubah = keterangan, True
            if matra and matra != ada[nama].matra:
                ada[nama].matra, berubah = matra, True
            if berubah:
                ubah.append(ada[nama])
            else:
                hasil["sama"] += 1
        else:
            baru[nama] = Satuan(nama=nama, keterangan=keterangan, matra=matra or matra_bawaan)
    Satuan.objects.bulk_create(baru.values(), batch_size=500)
    Satuan.objects.bulk_update(ubah, ["keterangan", "matra"], batch_size=500)
    hasil["baru"], hasil["diperbarui"] = len(baru), len(ubah)
    return hasil


@login_required
def satuan_impor(request):
    """GET: the upload form for the Pengaturan modal. POST: read the sheet."""
    from .views import _ke_tab_satuan

    if request.method != "POST":
        return render(request, "accounts/_satuan_impor.html")
    berkas = request.FILES.get("berkas")
    if berkas is None or not berkas.name.lower().endswith((".xlsx", ".csv")):
        messages.error(request, "Pilih berkas .xlsx atau .csv.")
        return _ke_tab_satuan(request)
    try:
        hasil = impor_baris(_baris_berkas(berkas))
    except Exception:  # a sheet that isn't one: openpyxl and csv raise all sorts
        messages.error(request, f"Berkas {berkas.name} tidak bisa dibaca.")
        return _ke_tab_satuan(request)
    messages.success(
        request,
        f"Import selesai: {hasil['baru']} satuan baru, {hasil['diperbarui']} diperbarui, "
        f"{hasil['sama']} sudah ada, {hasil['dilewati']} dilewati.",
    )
    return _ke_tab_satuan(request)
