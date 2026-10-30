"""The QR codes of a report's vehicles as a Word document, to print and cut."""
import io

from django.http import HttpResponse
from django.utils import timezone
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt

from fleet.qr import qr_png_with_logo

KOLOM = 3
# Six rows to an A4 page (the first one under the title too): a row exactly
# TINGGI_BARIS high holds the QR image (1.15 times as high as wide, with the
# No. Registrasi under the code) and the Merk/Type line.
BARIS_PER_HALAMAN = 6
TINGGI_BARIS = Mm(41)
LEBAR_QR = Mm(30)


def _rapat(paragraf):
    """No space around the paragraph: the rows' height is all for the code."""
    paragraf.paragraph_format.space_before = Pt(0)
    paragraf.paragraph_format.space_after = Pt(0)
    paragraf.paragraph_format.line_spacing = 1


def _tidak_terbelah(baris):
    """A row is never split across two pages."""
    properti = baris._tr.get_or_add_trPr()
    elemen = OxmlElement("w:cantSplit")
    elemen.set(qn("w:val"), "true")
    properti.append(elemen)


def qr_word(judul, kendaraan, nama_file):
    """A4 sheet(s) with a grid of ``kendaraan``'s QR codes (the same image as
    on the vehicle list: Kemhan logo, No. Registrasi underneath), each over
    its Merk/Type: KOLOM across, BARIS_PER_HALAMAN rows to a page."""
    doc = Document()
    bagian = doc.sections[0]
    bagian.page_width, bagian.page_height = Mm(210), Mm(297)
    bagian.left_margin = bagian.right_margin = Mm(15)
    bagian.top_margin = bagian.bottom_margin = Mm(15)

    kepala = doc.add_paragraph()
    _rapat(kepala)
    kepala.alignment = WD_ALIGN_PARAGRAPH.CENTER
    teks = kepala.add_run(f"QR Code — {judul}")
    teks.bold = True
    teks.font.size = Pt(14)
    ket = doc.add_paragraph()
    _rapat(ket)
    ket.paragraph_format.space_after = Pt(4)
    ket.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ket.add_run(f"{len(kendaraan)} kendaraan · dicetak {timezone.localtime():%d-%m-%Y %H:%M}").font.size = Pt(9)

    if not kendaraan:
        doc.add_paragraph("Tidak ada data untuk filter ini.").alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        tabel = doc.add_table(rows=-(-len(kendaraan) // KOLOM), cols=KOLOM)
        tabel.style = "Table Grid"
        tabel.alignment = WD_TABLE_ALIGNMENT.CENTER
        for baris in tabel.rows:
            baris.height = TINGGI_BARIS
            baris.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
            _tidak_terbelah(baris)
        for i, v in enumerate(kendaraan):
            sel = tabel.cell(i // KOLOM, i % KOLOM)
            sel.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            gambar = sel.paragraphs[0]
            _rapat(gambar)
            gambar.alignment = WD_ALIGN_PARAGRAPH.CENTER
            gambar.add_run().add_picture(io.BytesIO(qr_png_with_logo(v.qr_text(), label=v.nopol)), width=LEBAR_QR)
            bawah = sel.add_paragraph()
            _rapat(bawah)
            bawah.alignment = WD_ALIGN_PARAGRAPH.CENTER
            bawah.add_run(v.merek[:40]).font.size = Pt(8)

    isi = io.BytesIO()
    doc.save(isi)
    response = HttpResponse(
        isi.getvalue(), content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    response["Content-Disposition"] = f'attachment; filename="{nama_file}"'
    return response
